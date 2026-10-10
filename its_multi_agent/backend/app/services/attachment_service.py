"""按用户隔离存储聊天附件，并准备模型所需的附件内容。"""
import base64
import asyncio
import hashlib
import json
import mimetypes
import re
import uuid
import os
import subprocess
import tempfile
import zipfile
import html
from pathlib import Path
from typing import Any, Dict, List

from fastapi import UploadFile

from its_multi_agent.backend.app.config.settings import settings

MAX_CHAT_ATTACHMENT_BYTES = 10 * 1024 * 1024
ALLOWED_CHAT_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg", ".webp"}
MAX_CONTEXT_ATTACHMENTS = 8
MAX_CONTEXT_ATTACHMENT_CHARS = 60000


class AttachmentError(ValueError):
    pass


class AttachmentService:
    def __init__(self) -> None:
        base = Path(__file__).resolve().parent.parent / "data" / "chat_attachments"
        self._base = Path(getattr(settings, "CHAT_ATTACHMENT_FOLDER", None) or base)
        self._base.mkdir(parents=True, exist_ok=True)

    async def save(self, user_id: str, upload: UploadFile) -> Dict[str, Any]:
        name = Path(upload.filename or "").name
        ext = Path(name).suffix.lower()
        if not name or ext not in ALLOWED_CHAT_EXTENSIONS:
            raise AttachmentError("仅支持 txt、md、pdf、docx、xlsx、png、jpg、jpeg、webp 文件")
        content = await upload.read(MAX_CHAT_ATTACHMENT_BYTES + 1)
        if not content:
            raise AttachmentError("不能上传空文件")
        if len(content) > MAX_CHAT_ATTACHMENT_BYTES:
            raise AttachmentError("附件大小不能超过 10 MB")
        attachment_id = uuid.uuid4().hex
        folder = self._base / user_id
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{attachment_id}{ext}"
        await asyncio.to_thread(path.write_bytes, content)
        kind = "image" if ext in {".png", ".jpg", ".jpeg", ".webp"} else "document"
        item = {"id": attachment_id, "user_id": user_id, "name": name, "path": str(path), "kind": kind,
                "media_type": mimetypes.guess_type(name)[0] or "application/octet-stream", "size": len(content),
                "sha256": hashlib.sha256(content).hexdigest()}
        metadata = {key: value for key, value in item.items() if key != "path"}
        await asyncio.to_thread(
            path.with_suffix(".json").write_text, json.dumps(metadata, ensure_ascii=False), encoding="utf-8",
        )
        return item

    def _resolve(self, user_id: str, attachment_id: str):
        if not re.fullmatch(r"[0-9a-f]{32}", attachment_id):
            raise AttachmentError("无效的附件 ID")
        matches = [path for path in (self._base / user_id).glob(f"{attachment_id}.*")
                   if path.suffix.lower() in ALLOWED_CHAT_EXTENSIONS and path.is_file()]
        if len(matches) != 1:
            raise AttachmentError("附件不存在或无权访问")
        path = matches[0]
        metadata_path = path.with_suffix(".json")
        metadata = json.loads(metadata_path.read_text(encoding="utf-8")) if metadata_path.exists() else {}
        image = path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}
        return path, {
            "id": attachment_id, "name": metadata.get("name", path.name),
            "kind": "image" if image else "document",
            "media_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
        }

    def describe(self, user_id: str, attachment_ids: List[str]) -> List[Dict[str, Any]]:
        return [self._resolve(user_id, item)[1] for item in dict.fromkeys(attachment_ids)]

    def delete(self, user_id: str, attachment_id: str) -> bool:
        path, _metadata = self._resolve(user_id, attachment_id)
        path.unlink(missing_ok=True)
        path.with_suffix(".json").unlink(missing_ok=True)
        return True

    def prepare_messages(self, user_id: str, attachment_ids: List[str]) -> List[Dict[str, Any]]:
        messages: List[Dict[str, Any]] = []
        for attachment_id in attachment_ids[:8]:
            path, metadata = self._resolve(user_id, attachment_id)
            ext = path.suffix.lower()
            if ext in {".png", ".jpg", ".jpeg", ".webp"}:
                media_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
                data_url = f"data:{media_type};base64,{base64.b64encode(path.read_bytes()).decode('ascii')}"
                messages.append({"name": metadata["name"], "data_url": data_url})
            else:
                try:
                    text = self._extract_text(path, ext).strip()
                except AttachmentError:
                    raise
                except Exception as error:
                    raise AttachmentError(f"无法解析附件：{metadata['name']}，请检查文件是否完整。") from error
                if not text:
                    raise AttachmentError(f"无法提取文档内容：{metadata['name']}")
                messages.append({"name": metadata["name"], "text": text[:30000], "truncated": len(text) > 30000})
        return messages

    def prepare_history(
        self, user_id: str, history: List[Dict[str, Any]], *,
        prepared: Dict[str, Dict[str, Any]] | None = None, include_images: bool = True,
        evidence_out: List[Dict[str, Any]] | None = None,
    ) -> List[Dict[str, Any]]:
        # 优先使用最新证据，且不向模型 SDK 发送仅供归档使用的字段。
        selected = []
        for message in reversed(history):
            for reference in message.get("attachments", []):
                if reference["id"] not in selected and len(selected) < MAX_CONTEXT_ATTACHMENTS:
                    selected.append(reference["id"])
        available = dict(prepared or {})
        for item in selected:
            if item in available:
                continue
            try:
                _path, metadata = self._resolve(user_id, item)
                if metadata["kind"] == "image" and not include_images:
                    available[item] = {"name": metadata["name"], "text": "[当前模型无法读取历史图片，不能推断图片内容。]", "truncated": False, "unavailable": True}
                else:
                    available[item] = self.prepare_messages(user_id, [item])[0]
            except (AttachmentError, OSError, ValueError):
                # 历史证据丢失或损坏时，不能影响后续对话。
                continue
        prepared = {item: available[item] for item in selected if item in available}
        budget = MAX_CONTEXT_ATTACHMENT_CHARS
        for attachment in prepared.values():
            if "text" in attachment:
                full_text = attachment["text"]
                attachment["text"] = full_text[:budget]
                attachment["truncated"] = attachment["truncated"] or len(full_text) > budget
                budget -= len(attachment["text"])
        if evidence_out is not None:
            evidence_out.extend({
                "source_id": "attachment-" + item, "origin": "attachment",
                "source_name": attachment["name"], "excerpt": attachment["text"],
                "excerpt_truncated": attachment["truncated"],
            } for item, attachment in prepared.items() if attachment.get("text") and not attachment.get("unavailable"))
        result = []
        included = set()
        for message in history:
            content = message.get("content", "")
            parts = [{"type": "input_text", "text": content}]
            for reference in message.get("attachments", []):
                attachment = prepared.get(reference["id"])
                if attachment is None:
                    parts.append({"type": "input_text", "text": f"[附件 {reference['name']} 本轮无法加载，不能推断其内容。]"})
                elif reference["id"] in included:
                    parts.append({"type": "input_text", "text": f"[再次引用附件：{attachment['name']}]"})
                elif "text" in attachment:
                    included.add(reference["id"])
                    note = "\n[内容已截断；缺失部分无法据此确认。]" if attachment["truncated"] else ""
                    parts.append({"type": "input_text", "text": f"[附件：{attachment['name']}]\n{attachment['text']}{note}"})
                else:
                    included.add(reference["id"])
                    parts.append({"type": "input_image", "image_url": attachment["data_url"]})
            result.append({"role": message["role"], "content": parts if len(parts) > 1 else content})
        return result

    @staticmethod
    def _extract_text(path: Path, ext: str) -> str:
        if ext in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="replace")
        if ext == ".pdf":
            try:
                from pypdf import PdfReader
                return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
            except ImportError as error:
                raise AttachmentError("服务器未安装 PDF 文本提取组件") from error
        if ext == ".docx":
            try:
                from docx import Document
                doc = Document(str(path))
                from docx.table import Table
                from docx.text.paragraph import Paragraph
                blocks = []
                for element in doc.element.body:
                    if element.tag.endswith("}p"):
                        text = Paragraph(element, doc).text.strip()
                        if text:
                            blocks.append(text)
                        # python-docx 会遗漏绘图和文本框内的 w:t 节点。
                        xml_text = "".join(node.text or "" for node in element.iter()
                                           if node.tag.endswith("}t")).strip()
                        if xml_text and xml_text != text:
                            blocks.append(xml_text)
                    elif element.tag.endswith("}tbl"):
                        for row in Table(element, doc).rows:
                            row_text = " | ".join(cell.text.strip() for cell in row.cells).strip()
                            if row_text:
                                blocks.append(row_text)
                # 页眉、页脚和文本框存储在正文之外。
                for section in doc.sections:
                    for container in (section.header, section.first_page_header,
                                      section.even_page_header, section.footer,
                                      section.first_page_footer, section.even_page_footer):
                        for paragraph in container.paragraphs:
                            if paragraph.text.strip():
                                blocks.append(paragraph.text.strip())
                        for table in container.tables:
                            for row in table.rows:
                                row_text = " | ".join(cell.text.strip() for cell in row.cells).strip()
                                if row_text:
                                    blocks.append(row_text)
                with zipfile.ZipFile(path) as archive:
                    for name in archive.namelist():
                        if name.startswith("word/") and name.endswith(".xml"):
                            xml = archive.read(name).decode("utf-8", errors="ignore")
                            xml_text = " ".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml, re.S))
                            xml_text = re.sub(r"<[^>]+>", "", xml_text)
                            xml_text = html.unescape(xml_text).strip()
                            if xml_text and xml_text not in "\n".join(blocks):
                                blocks.append(xml_text)
                    ocr_blocks = []
                    for image_name in archive.namelist():
                        if image_name.startswith("word/media/"):
                            result = AttachmentService._ocr_bytes(archive.read(image_name), image_name)
                            if result:
                                ocr_blocks.append(f"图片 OCR：{result}")
                blocks.extend(ocr_blocks)
                # 尽量保持文档顺序，同时移除重复的页眉、XML 和 OCR 片段。
                return "\n".join(dict.fromkeys(item for item in blocks if item.strip()))
            except ImportError as error:
                raise AttachmentError("服务器未安装 Word 文本提取组件") from error
        if ext == ".xlsx":
            try:
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=True)
                try:
                    return "\n".join(f"[工作表：{sheet.title}]\n" + "\n".join(
                        " | ".join("" if value is None else str(value) for value in row)
                        for row in sheet.iter_rows(values_only=True)
                    ) for sheet in book.worksheets)
                finally:
                    book.close()
            except ImportError as error:
                raise AttachmentError("服务器未安装 Excel 文本提取组件") from error
        return ""

    @staticmethod
    def _ocr_bytes(data: bytes, name: str) -> str:
        """使用已安装的 Tesseract 对嵌入图片执行 OCR；不可用时返回空结果。"""
        command = os.getenv("TESSERACT_CMD", "tesseract")
        if command == "tesseract":
            candidates = [
                Path(os.environ.get("ProgramFiles", "C:\\Program Files")) / "Tesseract-OCR" / "tesseract.exe",
                Path("D:/dimage/python/software/Tesseract OCR/tesseract.exe"),
            ]
            command = next((str(item) for item in candidates if item.exists()), command)
        try:
            with tempfile.TemporaryDirectory(prefix="its-ocr-") as folder:
                source = Path(folder) / Path(name).name
                source.write_bytes(data)
                result = subprocess.run(
                    [command, str(source), "stdout", "-l", os.getenv("TESSERACT_LANG", "chi_sim+eng"), "--psm", "6"],
                    capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30,
                    check=False,
                )
                if result.returncode != 0:
                    return ""
                return re.sub(r"\s+", " ", result.stdout).strip()
        except (OSError, subprocess.SubprocessError):
            return ""


attachment_service = AttachmentService()
