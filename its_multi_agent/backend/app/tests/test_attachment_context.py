import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch
from types import SimpleNamespace

from fastapi import UploadFile

from its_multi_agent.backend.app.repositories.sqlite_session_repository import SQLiteSessionRepository
from its_multi_agent.backend.app.services.attachment_service import AttachmentError, AttachmentService
from its_multi_agent.backend.app.services.context_builder import ContextBuilder
from its_multi_agent.backend.app.services.session_service import SessionService


class NoCache:
    def get(self, *_args):
        return None

    def set(self, *_args):
        pass


class AttachmentContextTest(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.attachments = AttachmentService()
        self.attachments._base = Path(self.temp.name) / "attachments"
        self.attachments._base.mkdir()
        with patch.object(SQLiteSessionRepository, "_migrate_legacy_json"):
            self.repo = SQLiteSessionRepository(str(Path(self.temp.name) / "sessions.db"))

    async def upload(self, name="manual.txt", text="Warranty period: 36 months."):
        return await self.attachments.save("alice", UploadFile(file=io.BytesIO(text.encode()), filename=name))

    async def test_extracted_document_keeps_original_filename(self):
        item = await self.upload()
        prepared = self.attachments.prepare_messages("alice", [item["id"]])
        self.assertEqual(prepared[0]["name"], "manual.txt")
        self.assertIn("36 months", prepared[0]["text"])

    async def test_other_user_cannot_read_attachment(self):
        item = await self.upload()
        with self.assertRaises(AttachmentError):
            self.attachments.prepare_messages("bob", [item["id"]])

    async def test_archive_retains_attachment_references_without_file_data(self):
        item = await self.upload()
        reference = {key: item[key] for key in ("id", "name", "kind", "media_type")}
        self.repo.append_message("alice", "session", {
            "role": "user", "content": "Read this", "attachments": [reference],
        })
        archived = self.repo.load_session("alice", "session")
        self.assertEqual(archived[0].get("attachments"), [reference])
        self.assertNotIn("36 months", json.dumps(archived))
        self.assertNotIn(str(self.attachments._base), json.dumps(archived))

    async def test_identical_question_with_different_attachment_is_not_dropped(self):
        for name in ("first.txt", "second.txt"):
            item = await self.upload(name)
            self.repo.append_message("alice", "session", {
                "role": "user", "content": "Read this", "attachments": [{"id": item["id"], "name": name}],
            })
        self.assertEqual(len(self.repo.load_session("alice", "session")), 2)

    async def test_follow_up_rehydrates_archived_evidence(self):
        item = await self.upload()
        reference = {key: item[key] for key in ("id", "name", "kind", "media_type")}
        service = SessionService()
        service._repo = self.repo
        builder = ContextBuilder(self.repo, NoCache())
        with patch("its_multi_agent.backend.app.services.session_service.context_builder", builder):
            history = service.prepare_history("alice", "session", "Read this", attachments=[reference])
            history.append({"role": "assistant", "content": "The document is available."})
            service.save_history("alice", "session", history)
            following = service.prepare_history("alice", "session", "How long is the warranty?")
        inputs = self.attachments.prepare_history("alice", following)
        self.assertIn("36 months", json.dumps(inputs))
        self.assertTrue(all("attachments" not in message for message in inputs))
        self.assertEqual(following[0]["role"], "system")

    async def test_missing_historical_file_does_not_block_new_question(self):
        item = await self.upload()
        Path(item["path"]).unlink()
        history = [{"role": "user", "content": "Read this", "attachments": [{"id": item["id"], "name": item["name"]}]},
                   {"role": "user", "content": "Hello"}]
        prepared = self.attachments.prepare_history("alice", history)
        self.assertEqual(prepared[-1]["content"], "Hello")
        self.assertIn("无法", json.dumps(prepared, ensure_ascii=False))

    async def test_text_model_can_continue_after_an_old_image(self):
        item = await self.attachments.save("alice", UploadFile(file=io.BytesIO(b"image"), filename="diagram.png"))
        history = [{"role": "user", "content": "Image", "attachments": [{"id": item["id"], "name": item["name"]}]},
                   {"role": "user", "content": "Hello"}]
        prepared = self.attachments.prepare_history("alice", history, include_images=False)
        self.assertNotIn("input_image", json.dumps(prepared))
        self.assertIn("图片", json.dumps(prepared, ensure_ascii=False))

    async def test_hydration_exposes_bounded_document_evidence_for_review(self):
        item = await self.upload()
        evidence = []
        history = [{"role": "user", "content": "Read this", "attachments": [
            {"id": item["id"], "name": item["name"]},
        ]}]
        self.attachments.prepare_history("alice", history, evidence_out=evidence)
        self.assertEqual(evidence[0]["source_id"], "attachment-" + item["id"])
        self.assertIn("36 months", evidence[0]["excerpt"])
        self.assertNotIn("data_url", evidence[0])

    async def test_corrupt_current_file_is_rejected_before_archival(self):
        from its_multi_agent.backend.app.schemas.request import ChatMessageRequest, UserContext
        from its_multi_agent.backend.app.services.agent_service import MultiAgentService
        item = await self.attachments.save("alice", UploadFile(file=io.BytesIO(b"not a zip"), filename="broken.docx"))
        target = "its_multi_agent.backend.app.services.agent_service."
        with (
            patch(target + "attachment_service", self.attachments),
            patch(target + "memory_service.extract_and_store", AsyncMock()),
            patch(target + "resolve_profile", return_value=SimpleNamespace(supports_vision=False, supports_tools=False)),
            patch(target + "model_for", return_value=object()),
            patch(target + "build_orchestrator_agent", return_value=object()),
            patch(target + "session_service.prepare_history") as archive,
        ):
            packets = [json.loads(packet.removeprefix("data: ")) async for packet in MultiAgentService.process_task(
                ChatMessageRequest(query="Read", context=UserContext(session_id="s"), attachment_ids=[item["id"]]),
                user_id="alice", flag=False,
            )]
        archive.assert_not_called()
        self.assertEqual(packets[-1]["metadata"]["finishReason"], "ERROR")
        self.assertIn("broken.docx", packets[0]["content"]["text"])

    def test_docx_table_text_is_included(self):
        from docx import Document
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Warranty"
        table.cell(0, 1).text = "36 months"
        path = Path(self.temp.name) / "table.docx"
        document.save(path)
        self.assertIn("36 months", self.attachments._extract_text(path, ".docx"))


if __name__ == "__main__":
    unittest.main()
