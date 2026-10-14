import logging
import re
import unicodedata
from io import BytesIO
from pathlib import Path
from threading import Lock
from typing import List

from langchain_community.document_loaders import TextLoader
from langchain_core.documents import Document

from its_multi_agent.backend.knowledge.config.settings import settings


logger = logging.getLogger(__name__)


class DocumentExtractor:
    """从支持的知识库文档中提取文本和 OCR 内容。"""

    _ocr_engine = None
    _ocr_lock = Lock()

    @classmethod
    def load(cls, file_path: str) -> List[Document]:
        extension = Path(file_path).suffix.lower()
        if extension == ".md":
            return cls._load_markdown(file_path)
        if extension == ".txt":
            documents = TextLoader(
                file_path=file_path, encoding="utf-8", autodetect_encoding=True
            ).load()
            return [cls._clean_document(document) for document in documents]
        if extension == ".pdf":
            return cls._load_pdf(file_path)
        if extension == ".docx":
            return cls._load_docx(file_path)
        if extension == ".xlsx":
            return cls._load_xlsx(file_path)
        raise ValueError(f"不支持的文档格式: {extension}")

    @staticmethod
    def _load_markdown(file_path: str) -> List[Document]:
        loaded = TextLoader(
            file_path=file_path,
            encoding="utf-8",
            autodetect_encoding=True,
        ).load()
        text = "\n\n".join(document.page_content for document in loaded)
        base_metadata = loaded[0].metadata if loaded else {"source": file_path}
        documents = []
        headers = {}
        section_lines = []
        has_body = False
        last_heading_level = 0
        fence = ""
        heading_pattern = re.compile(r"^ {0,3}(#{1,6})(?:[ \t]+(.*)|[ \t]*)$")
        fence_pattern = re.compile(r"^ {0,3}(`{3,}|~{3,})(.*)$")
        closing_fence_pattern = re.compile(r"^ {0,3}(`{3,}|~{3,})[ \t]*$")

        def append_section():
            content = "".join(section_lines)
            if not content.strip():
                section_lines.clear()
                return
            metadata = {**base_metadata, **headers}
            section_path = " / ".join(
                headers[f"h{level}"]
                for level in range(1, 7)
                if headers.get(f"h{level}")
            )
            if section_path:
                metadata["section_path"] = section_path
            documents.append(Document(page_content=content, metadata=metadata))
            section_lines.clear()

        # 只检测标题和代码围栏，保留原始行的缩进、制表符及换行。
        for line in text.splitlines(keepends=True):
            plain_line = line.rstrip("\r\n")
            if fence:
                section_lines.append(line)
                closing = closing_fence_pattern.match(plain_line)
                if (
                    closing
                    and closing.group(1)[0] == fence[0]
                    and len(closing.group(1)) >= len(fence)
                ):
                    fence = ""
                continue
            opening = fence_pattern.match(plain_line)
            if opening and not (opening.group(1)[0] == "`" and "`" in opening.group(2)):
                fence = opening.group(1)
                section_lines.append(line)
                has_body = True
                continue
            heading = heading_pattern.match(plain_line)
            if heading:
                level = len(heading.group(1))
                if has_body or (last_heading_level and level <= last_heading_level):
                    append_section()
                    has_body = False
                for header_level in range(level, 7):
                    headers.pop(f"h{header_level}", None)
                title = re.sub(r"[ \t]+#+[ \t]*$", "", heading.group(2) or "").strip()
                headers[f"h{level}"] = title
                last_heading_level = level
            elif plain_line.strip():
                has_body = True
            section_lines.append(line)
        append_section()
        return documents

    @classmethod
    def _load_pdf(cls, file_path: str) -> List[Document]:
        from pypdf import PdfReader

        reader = PdfReader(file_path)
        pdfium_document = None
        documents = []
        try:
            for page_index, page in enumerate(reader.pages):
                text = (page.extract_text() or "").strip()
                ocr_text = ""
                if settings.OCR_ENABLED and len(text) < settings.OCR_PDF_TEXT_THRESHOLD:
                    if pdfium_document is None:
                        import pypdfium2 as pdfium

                        pdfium_document = pdfium.PdfDocument(file_path)
                    scale = settings.OCR_PDF_DPI / 72
                    image = pdfium_document[page_index].render(scale=scale).to_pil()
                    ocr_text = cls._recognize_image(image)
                page_content = cls._join_unique_text(text, ocr_text)
                page_content = cls._clean_text(page_content)
                documents.append(
                    Document(
                        page_content=page_content,
                        metadata={
                            "page": page_index + 1,
                            "ocr_applied": bool(ocr_text),
                            "extraction_method": (
                                "text+ocr" if text and ocr_text else
                                "ocr" if ocr_text else
                                "text" if text else "empty"
                            ),
                        },
                    )
                )
        finally:
            if pdfium_document is not None:
                pdfium_document.close()
        return documents

    @classmethod
    def _load_docx(cls, file_path: str) -> List[Document]:
        from docx import Document as WordDocument
        from docx.oxml.ns import qn
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        source = WordDocument(file_path)
        image_index = 0

        def paragraph_text(paragraph) -> str:
            nonlocal image_index
            parts = []
            # XPath 联集保持普通 run 与超链接内 run 的文档顺序。
            for run_element in paragraph._p.xpath("./w:r | ./w:hyperlink/w:r"):
                for item in run_element.xpath(
                    "w:br | w:cr | w:noBreakHyphen | w:ptab | w:t | w:tab | w:drawing"
                ):
                    if item.tag != qn("w:drawing"):
                        parts.append(str(item))
                        continue
                    if not settings.OCR_ENABLED:
                        continue
                    for image in item.xpath(".//a:blip"):
                        relationship_id = image.get(qn("r:embed"))
                        image_part = source.part.related_parts.get(relationship_id)
                        if image_part is None:
                            continue
                        image_index += 1
                        ocr_text = cls._recognize_image(image_part.blob)
                        if ocr_text:
                            parts.append(f"\n图片 {image_index} OCR: {ocr_text}\n")
            return "".join(parts).strip()

        def render_blocks(container, table_counter) -> list[str]:
            element = (
                container.element.body
                if hasattr(container, "element")
                else container._tc
            )
            blocks = []
            for child in element.iterchildren():
                if isinstance(child, CT_P):
                    text = paragraph_text(Paragraph(child, container))
                    if text:
                        blocks.append(text)
                elif isinstance(child, CT_Tbl):
                    table_counter[0] += 1
                    table_number = table_counter[0]
                    table = Table(child, container)
                    rows = []
                    for row in table.rows:
                        cells = [
                            "\n".join(render_blocks(cell, table_counter))
                            for cell in row.cells
                        ]
                        rendered_row = " | ".join(value.strip() for value in cells)
                        if rendered_row.strip(" |\n"):
                            rows.append(rendered_row)
                    if rows:
                        blocks.append(f"表格 {table_number}:\n" + "\n".join(rows))
            return blocks

        blocks = render_blocks(source, [0])
        content = cls._clean_text("\n\n".join(blocks))
        return [Document(
            page_content=content,
            metadata={"extraction_method": "docx"},
        )]

    @staticmethod
    def _load_xlsx(file_path: str) -> List[Document]:
        from openpyxl import load_workbook

        formula_workbook = load_workbook(file_path, read_only=True, data_only=False)
        value_workbook = load_workbook(file_path, read_only=True, data_only=True)
        documents = []
        try:
            for sheet in formula_workbook.worksheets:
                value_sheet = value_workbook[sheet.title]
                rows = []
                formula_rows = sheet.iter_rows(values_only=False)
                value_rows = value_sheet.iter_rows(values_only=False)
                for formula_row, value_row in zip(formula_rows, value_rows):
                    values = []
                    for cell, value_cell in zip(formula_row, value_row):
                        value = cell.value
                        if value is None or not str(value).strip():
                            continue
                        if cell.data_type == "f":
                            cached_value = value_cell.value
                            rendered = (
                                f"{value} (计算值: {cached_value})"
                                if cached_value is not None
                                else str(value)
                            )
                        else:
                            rendered = str(value).strip()
                        values.append(f"{cell.coordinate}: {rendered}")
                    if values:
                        rows.append(" | ".join(values))
                if rows:
                    documents.append(
                        Document(
                            page_content=cls._clean_text(
                                f"工作表: {sheet.title}\n" + "\n".join(rows)
                            ),
                            metadata={
                                "sheet": sheet.title,
                                "sheet_state": sheet.sheet_state,
                                "extraction_method": "xlsx-cells",
                            },
                        )
                    )
        finally:
            formula_workbook.close()
            value_workbook.close()
        return documents

    @classmethod
    def _recognize_image(cls, image) -> str:
        try:
            from PIL import Image
            import numpy as np

            if isinstance(image, bytes):
                image = Image.open(BytesIO(image)).convert("RGB")
            else:
                image = image.convert("RGB")
            result, _ = cls._get_ocr_engine()(np.asarray(image))
            if not result:
                return ""
            lines = [
                str(item[1]).strip()
                for item in result
                if len(item) > 1 and str(item[1]).strip()
            ]
            return "\n".join(lines)
        except Exception as error:
            logger.warning("图片 OCR 失败: %s", error)
            return ""

    @classmethod
    def _get_ocr_engine(cls):
        if cls._ocr_engine is None:
            with cls._ocr_lock:
                if cls._ocr_engine is None:
                    from rapidocr_onnxruntime import RapidOCR

                    cls._ocr_engine = RapidOCR()
        return cls._ocr_engine

    @staticmethod
    def _join_unique_text(text: str, ocr_text: str) -> str:
        if not text:
            return ocr_text
        if not ocr_text or DocumentExtractor._is_text_present(text, ocr_text):
            return text
        return f"{text}\n\nOCR 提取内容:\n{ocr_text}"

    @staticmethod
    def _is_text_present(original: str, candidate: str) -> bool:
        """判断 OCR 结果是否已经存在，兼容换行和空白差异。"""
        compact_original = re.sub(r"\s+", "", original)
        compact_candidate = re.sub(r"\s+", "", candidate)
        return bool(compact_candidate) and compact_candidate in compact_original

    @staticmethod
    def _clean_text(value: str) -> str:
        """清理跨格式抽取噪声，同时保留表格和代码中的有效换行。"""
        if not value:
            return ""
        normalized = unicodedata.normalize("NFC", str(value))
        normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")
        # 去掉不可见控制字符和零宽字符；制表符、换行和普通空格保留。
        normalized = "".join(
            char for char in normalized
            if char in {"\n", "\t"} or not unicodedata.category(char).startswith("C")
        )
        lines = [line.rstrip() for line in normalized.split("\n")]
        cleaned_lines = []
        blank_count = 0
        for line in lines:
            if line.strip():
                blank_count = 0
                cleaned_lines.append(line)
            elif blank_count < 2:
                blank_count += 1
                cleaned_lines.append("")
        return "\n".join(cleaned_lines).strip()

    @classmethod
    def _clean_document(cls, document: Document) -> Document:
        """清洗普通文本文档，保留 LangChain metadata。"""
        return Document(
            page_content=cls._clean_text(document.page_content),
            metadata=dict(document.metadata),
        )
