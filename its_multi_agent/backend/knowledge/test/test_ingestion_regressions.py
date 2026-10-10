import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from its_multi_agent.backend.knowledge.batch_upload import upload_file, wait_for_task
from its_multi_agent.backend.knowledge.config.settings import settings
from its_multi_agent.backend.knowledge.services.crawler.client import KnowledgeApiClient
from its_multi_agent.backend.knowledge.services.crawler.parser import HtmlParser
from its_multi_agent.backend.knowledge.services.ingestion.document_extractor import (
    DocumentExtractor,
)
from its_multi_agent.backend.knowledge.utils.markdown_utils import MarkDownUtils


class Response:
    def __init__(self, status_code, payload):
        self.status_code = status_code
        self._payload = payload

    def json(self):
        return self._payload

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


class IngestionRegressionTest(unittest.TestCase):
    def test_markdown_image_cleaner_is_callable_and_keeps_urls(self):
        text = '说明 ![截图](https://example.test/image.png "标题") 结束'
        self.assertEqual(
            MarkDownUtils.clean_markdown_images(text),
            "说明 \nhttps://example.test/image.png\n 结束",
        )

    def test_crawler_client_wraps_requests_errors(self):
        with patch(
            "its_multi_agent.backend.knowledge.services.crawler.client.settings.KNOWLEDGE_BASE_URL",
            "https://example.test",
        ), patch(
            "its_multi_agent.backend.knowledge.services.crawler.client.requests.get",
            side_effect=__import__("requests").RequestException("offline"),
        ):
            with self.assertRaisesRegex(RuntimeError, "远端知识库失败"):
                KnowledgeApiClient.fetch_knowledge_content("1")

    def test_crawler_parser_accepts_optional_fields(self):
        markdown = HtmlParser().parse_html_to_markdown(
            "1", {"title": "标题", "content": "<p>答案</p>"}
        )
        self.assertIn("## 标题", markdown)
        self.assertIn("## 解决方案", markdown)
        self.assertIn("答案", markdown)

    def test_crawler_parser_rejects_content_empty_after_cleanup(self):
        for content in (
            "<p>&nbsp;</p>",
            "<script>alert(1)</script><style>p { color: red; }</style>",
        ):
            with self.subTest(content=content), self.assertRaises(ValueError):
                HtmlParser().parse_html_to_markdown("1", {"content": content})

    def test_batch_upload_waits_for_async_task(self):
        with tempfile.TemporaryDirectory() as directory:
            file_path = Path(directory) / "one.md"
            file_path.write_text("内容", encoding="utf-8")
            session = Mock()
            session.post.return_value = Response(
                202, {"status": "queued", "task_id": "task-1"}
            )
            session.get.side_effect = [
                Response(200, {"status": "processing"}),
                Response(200, {"status": "success", "chunks_added": 2}),
            ]
            clock = iter([0.0, 0.0, 1.0, 2.0])
            self.assertEqual(
                upload_file(
                    session,
                    file_path,
                    "http://127.0.0.1:8001",
                    knowledge_space="default",
                    domain="general",
                    poll_interval=0.1,
                    task_timeout=10,
                    attempts=1,
                    retry_delay=0,
                    sleep=lambda _: None,
                    monotonic=lambda: next(clock, 2.0),
                ),
                (True, "入库完成"),
            )
            self.assertEqual(session.get.call_count, 2)

    def test_batch_upload_returns_failure_when_task_polling_fails(self):
        import requests

        with tempfile.TemporaryDirectory() as directory:
            file_path = Path(directory) / "one.md"
            file_path.write_text("内容", encoding="utf-8")
            session = Mock()
            session.post.return_value = Response(
                202, {"status": "queued", "task_id": "task-1"}
            )
            session.get.side_effect = requests.ConnectionError("offline")

            ok, message = upload_file(
                session,
                file_path,
                "http://127.0.0.1:8001",
                knowledge_space="default",
                domain="general",
                poll_interval=0.1,
                task_timeout=10,
                attempts=2,
                retry_delay=0,
                sleep=lambda _: None,
            )

        self.assertFalse(ok)
        self.assertIn("查询任务请求失败", message)
        self.assertEqual(session.get.call_count, 2)

    def test_task_polling_stops_retrying_at_deadline(self):
        import requests

        clock = [0.0]

        def sleep(seconds):
            clock[0] += seconds

        session = Mock()
        session.get.side_effect = requests.ConnectionError("offline")
        ok, message = wait_for_task(
            session,
            "http://127.0.0.1:8001",
            "task-1",
            poll_interval=1,
            task_timeout=5,
            attempts=3,
            retry_delay=10,
            sleep=sleep,
            monotonic=lambda: clock[0],
        )
        self.assertFalse(ok)
        self.assertIn("等待入库超时", message)
        self.assertEqual(session.get.call_count, 1)
        self.assertEqual(clock[0], 5)

    def test_markdown_extractor_keeps_heading_path(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "guide.md"
            path.write_text("# 安装\n\n## Windows\n步骤一", encoding="utf-8")
            documents = DocumentExtractor.load(str(path))
        self.assertEqual(len(documents), 1)
        self.assertEqual(documents[0].metadata["section_path"], "安装 / Windows")
        self.assertIn("步骤一", documents[0].page_content)

    def test_markdown_sections_preserve_code_and_list_indentation(self):
        text = (
            "# 手册\n\n## Python\n```python\nif ok:\n    reboot()\n\tlog()\n"
            "# 代码内部标题\n```\n\n- 第一步\n  - 第二步\n\n"
            "## YAML\n~~~yaml\nservice:\n  enabled: true\n~~~\n\n"
            "    if ready:\n        run()\n"
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "guide.md"
            path.write_text(text, encoding="utf-8")
            documents = DocumentExtractor.load(str(path))
        self.assertEqual("".join(document.page_content for document in documents), text)
        self.assertEqual(
            [document.metadata["section_path"] for document in documents],
            ["手册 / Python", "手册 / YAML"],
        )

    def test_docx_extractor_preserves_paragraph_and_table_order(self):
        from docx import Document as WordDocument

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "guide.docx"
            source = WordDocument()
            source.add_paragraph("表格前")
            table = source.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "名称"
            table.rows[0].cells[1].text = "值"
            source.add_paragraph("表格后")
            source.save(path)
            with patch.object(settings, "OCR_ENABLED", False):
                content = DocumentExtractor.load(str(path))[0].page_content
        self.assertLess(content.index("表格前"), content.index("表格 1"))
        self.assertLess(content.index("表格 1"), content.index("表格后"))
        self.assertIn("名称 | 值", content)

    def test_docx_extractor_runs_ocr_for_images_inside_tables(self):
        from docx import Document as WordDocument
        from PIL import Image

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image_path = root / "pixel.png"
            Image.new("RGB", (4, 4), "white").save(image_path)
            document_path = root / "table.docx"
            source = WordDocument()
            cell = source.add_table(rows=1, cols=1).cell(0, 0)
            cell.paragraphs[0].add_run().add_picture(str(image_path))
            source.save(document_path)
            with patch.object(settings, "OCR_ENABLED", True), patch.object(
                DocumentExtractor, "_recognize_image", return_value="表格图片文字"
            ) as recognize:
                content = DocumentExtractor.load(str(document_path))[0].page_content
        self.assertIn("表格图片文字", content)
        recognize.assert_called_once()

    def test_docx_extractor_keeps_hyperlink_text_in_paragraphs_and_tables(self):
        from docx import Document as WordDocument
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "links.docx"
            source = WordDocument()
            paragraph = source.add_paragraph("点击")
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "download")
            hyperlink.append(paragraph.add_run("下载驱动")._r)
            paragraph._p.append(hyperlink)
            paragraph.add_run("后继续")
            cell_paragraph = source.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
            cell_link = OxmlElement("w:hyperlink")
            cell_link.set(qn("w:anchor"), "details")
            cell_link.append(cell_paragraph.add_run("表格资料")._r)
            cell_paragraph._p.append(cell_link)
            source.save(path)
            with patch.object(settings, "OCR_ENABLED", False):
                content = DocumentExtractor.load(str(path))[0].page_content
        self.assertIn("点击下载驱动后继续", content)
        self.assertIn("表格资料", content)

    def test_docx_extractor_recognizes_images_inside_hyperlinks(self):
        from docx import Document as WordDocument
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from PIL import Image

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image_path = root / "pixel.png"
            Image.new("RGB", (4, 4), "white").save(image_path)
            source = WordDocument()
            paragraph = source.add_paragraph()
            run = paragraph.add_run()
            run.add_picture(str(image_path))
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "image")
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
            path = root / "image_link.docx"
            source.save(path)
            with patch.object(settings, "OCR_ENABLED", True), patch.object(
                DocumentExtractor, "_recognize_image", return_value="链接图片文字"
            ) as recognize:
                content = DocumentExtractor.load(str(path))[0].page_content
        self.assertIn("链接图片文字", content)
        recognize.assert_called_once()

    def test_docx_nested_tables_have_distinct_numbers(self):
        from docx import Document as WordDocument

        with tempfile.TemporaryDirectory() as directory:
            source = WordDocument()
            outer_cell = source.add_table(rows=1, cols=1).cell(0, 0)
            outer_cell.text = "外层"
            outer_cell.add_table(rows=1, cols=1).cell(0, 0).text = "内层"
            source.add_table(rows=1, cols=1).cell(0, 0).text = "下一个表格"
            path = Path(directory) / "nested.docx"
            source.save(path)
            with patch.object(settings, "OCR_ENABLED", False):
                content = DocumentExtractor.load(str(path))[0].page_content
        for number in (1, 2, 3):
            self.assertEqual(content.count(f"表格 {number}:"), 1)
        self.assertLess(content.index("表格 1:"), content.index("表格 2:"))
        self.assertLess(content.index("表格 2:"), content.index("表格 3:"))

    def test_docx_extractor_keeps_text_and_image_order_within_one_run(self):
        from docx import Document as WordDocument
        from PIL import Image

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image_path = root / "pixel.png"
            Image.new("RGB", (4, 4), "white").save(image_path)
            source = WordDocument()
            run = source.add_paragraph().add_run("图前\t")
            run.add_picture(str(image_path))
            run.add_text("\n图后")
            path = root / "inline_image.docx"
            source.save(path)
            with patch.object(settings, "OCR_ENABLED", True), patch.object(
                DocumentExtractor, "_recognize_image", return_value="图片文字"
            ) as recognize:
                content = DocumentExtractor.load(str(path))[0].page_content
            with patch.object(settings, "OCR_ENABLED", False):
                text_only = DocumentExtractor.load(str(path))[0].page_content
        self.assertLess(content.index("图前"), content.index("图片文字"))
        self.assertLess(content.index("图片文字"), content.index("图后"))
        self.assertEqual(text_only, "图前\t\n图后")
        recognize.assert_called_once()

    def test_xlsx_extractor_keeps_cell_coordinates_and_formulas(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "guide.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "设备"
            sheet["A1"] = "产品"
            sheet["B1"] = "数量"
            sheet["A2"] = "机型A"
            sheet["B2"] = "=1+1"
            workbook.save(path)
            documents = DocumentExtractor.load(str(path))
        self.assertEqual(documents[0].metadata["sheet"], "设备")
        self.assertIn("A2: 机型A", documents[0].page_content)
        self.assertIn("B2: =1+1", documents[0].page_content)

    def test_upload_timeout_then_existing_task_continues_polling(self):
        import requests
        with tempfile.TemporaryDirectory() as directory:
            document = Path(directory) / "file.md"
            document.write_text("content", encoding="utf-8")
            session = Mock()
            session.post.side_effect = [requests.ReadTimeout("response lost"), Response(202, {"status": "queued", "task_id": "existing"})]
            session.get.return_value = Response(200, {"status": "completed"})
            ok, _ = upload_file(session, document, "http://example.test", knowledge_space="default", domain="general",
                                poll_interval=0.1, task_timeout=5, attempts=2, retry_delay=0, sleep=lambda _: None)
            self.assertTrue(ok)
            self.assertIn("/tasks/existing", session.get.call_args.args[0])


if __name__ == "__main__":
    unittest.main()
